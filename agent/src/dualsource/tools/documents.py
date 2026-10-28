from __future__ import annotations

import os
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


@dataclass
class DocumentResult:
    status: str
    files: list[dict[str, str]]
    error: str | None = None


def _safe_mkdir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def generate_pptx(
    title: str,
    summary: str,
    bullets: list[str] | None = None,
    table: dict[str, Any] | None = None,
    output_dir: str | Path | None = None,
    user_email: str = "unknown",
    session_id: str = "session",
) -> DocumentResult:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception as exc:  # noqa: BLE001
        return DocumentResult(status="error", files=[], error=f"python-pptx missing: {exc}")

    out_root = Path(output_dir or os.getenv("CI_FILES_DIR") or "./tmp_files")
    target_dir = out_root / user_email / session_id
    _safe_mkdir(target_dir)
    file_id = uuid.uuid4().hex
    filename = f"{file_id}_dualsource.pptx"
    path = target_dir / filename

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title[:120] or "DualSource Agent"
    body = slide.placeholders[1].text_frame
    body.clear()
    p = body.paragraphs[0]
    p.text = summary[:1500]
    p.font.size = Pt(16)
    for b in (bullets or [])[:8]:
        para = body.add_paragraph()
        para.text = b[:300]
        para.level = 1

    if table and table.get("columns") and table.get("rows"):
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])
        slide2.shapes.title.text = table.get("title") or "Data"
        cols = list(table["columns"])[:6]
        rows = list(table["rows"])[:12]
        shape = slide2.shapes.add_table(len(rows) + 1, len(cols), Inches(0.5), Inches(1.5), Inches(9), Inches(4)).table
        for i, c in enumerate(cols):
            shape.cell(0, i).text = str(c)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(list(row)[: len(cols)]):
                shape.cell(r_idx, c_idx).text = str(val)

    notes = prs.slides.add_slide(prs.slide_layouts[1])
    notes.shapes.title.text = "Notes & Sources"
    tf = notes.placeholders[1].text_frame
    tf.text = f"Generated at {datetime.now(timezone.utc).isoformat()} UTC\nSource labels retained from DualSource Agent."

    prs.save(path)
    return DocumentResult(
        status="ok",
        files=[{
            "id": file_id,
            "type": "pptx",
            "filename": filename,
            "download_path": str(path),
        }],
    )


def _plain_summary(text: str, limit: int = 2000) -> str:
    """Strip markdown noise for Excel Summary sheet readability."""
    import re

    s = str(text or "")
    s = re.sub(r"\[\[CHART(?::\d+)?\]\]", " ", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)  # [label](url) -> label
    s = s.replace("**", "").replace("__", "").replace("`", "")
    s = s.replace("[Internal]", "Internal")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()[:limit]


def generate_xlsx(
    title: str,
    table: dict[str, Any] | None = None,
    summary: str = "",
    output_dir: str | Path | None = None,
    user_email: str = "unknown",
    session_id: str = "session",
) -> DocumentResult:
    try:
        from openpyxl import Workbook
    except Exception as exc:  # noqa: BLE001
        return DocumentResult(status="error", files=[], error=f"openpyxl missing: {exc}")

    out_root = Path(output_dir or os.getenv("CI_FILES_DIR") or "./tmp_files")
    target_dir = out_root / user_email / session_id
    _safe_mkdir(target_dir)
    file_id = uuid.uuid4().hex
    filename = f"{file_id}_dualsource.xlsx"
    path = target_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title or "DualSource Agent"
    ws["A2"] = _plain_summary(summary)
    ws["A4"] = f"Generated at {datetime.now(timezone.utc).isoformat()} UTC"

    if table and table.get("columns") and table.get("rows"):
        ws2 = wb.create_sheet("Data")
        cols = list(table["columns"])
        ws2.append(cols)
        for row in list(table["rows"])[:500]:
            ws2.append(list(row)[: len(cols)])

    wb.save(path)
    return DocumentResult(
        status="ok",
        files=[{
            "id": file_id,
            "type": "xlsx",
            "filename": filename,
            "download_path": str(path),
        }],
    )
